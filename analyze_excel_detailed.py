import openpyxl

wb = openpyxl.load_workbook(r'c:\Users\lexus\Documents\Projects\Fazenda_app_android\Електронна таблиця без назви.xlsx')
ws1 = wb['Лист1']

print(f"=== ЛИСТ 1: ПОВНА СТРУКТУРА ===")
print(f"Max row: {ws1.max_row}, Max col: {ws1.max_column}\n")

for row_idx in range(1, ws1.max_row + 1):
    col_a = ws1.cell(row=row_idx, column=1).value
    col_b = ws1.cell(row=row_idx, column=2).value
    
    if col_a or col_b:
        print(f"Рядок {row_idx:2}: A='{col_a}' | B='{col_b}'")

print("\n\n=== ВИНОГРАД У ЛИСТ2 ===")
ws2 = wb['Виноград (копия)']

print(f"Перших 40 рядків з даними:")
count = 0
for row_idx in range(2, ws2.max_row + 1):
    name = ws2.cell(row=row_idx, column=3).value
    if name:
        count += 1
        row_data = [
            ws2.cell(row=row_idx, column=1).value,
            ws2.cell(row=row_idx, column=2).value,
            name,
            ws2.cell(row=row_idx, column=4).value,
            ws2.cell(row=row_idx, column=5).value,
            ws2.cell(row=row_idx, column=6).value,
        ]
        if count <= 40:
            print(f"  Рядок {row_idx}: {row_data}")
        if count == 40:
            print(f"  ... (всього {ws2.max_row - 1} потенційних рядків)")
            break

# Count all grape rows
total_grapes = 0
for row_idx in range(2, ws2.max_row + 1):
    if ws2.cell(row=row_idx, column=3).value:
        total_grapes += 1

print(f"\nВсього записів про Виноград: {total_grapes}")
